import sys
import unittest
import datetime as dt
from pathlib import Path
from unittest.mock import Mock, patch

import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
from helpers import ROOT, Sandbox  # noqa: E402

sys.path.insert(0, str(ROOT / "scripts"))
import ingest  # noqa: E402
import schema  # noqa: E402
import workbook_spec as spec  # noqa: E402

TEMPLATE = ROOT / "templates" / "ingestion_template.xlsx"
EXAMPLE = ROOT / "templates" / "ingestion_example_filled.xlsx"
COLS = {"INGESTION": spec.INGESTION_COLUMNS, **spec.SIDE_COLUMNS}


def fill(path, sheets):
    """Write rows into the template (INGESTION: header on row 2, data from row 3; side sheets: header on row 1)."""
    wb = openpyxl.load_workbook(TEMPLATE)
    for name, rows in sheets.items():
        ws = wb[name]
        for r in rows:
            ws.append([r.get(c, "") for c in COLS[name]])
    wb.save(path)


class IngestBase(unittest.TestCase):
    def setUp(self):
        self.sb = Sandbox("live")
        ev = pd.read_csv(self.sb.data / "events.csv", dtype=str, keep_default_na=False)
        self.ev = ev[ev.event_name.str.contains("Ridgecrest 2")].event_id.iloc[0]
        self.pub = dict(citation_short="Testa et al. (2030)", citation_full="Testa, A. and Bianchi, B. (2030). A test paper. J. Test., 1, 1-10.",
                        doi="10.1000/test.2030.1")

    def tearDown(self):
        self.sb.cleanup()

    def row(self, **kw):
        base = dict(target_id=self.ev, metric_id="mapped_surface_rupture_length", source_id="Testa et al. (2030)", statistic_type="individual",
                    source_location="Table 1", method="field")
        base.update(kw)
        return base

    def ingest(self, name, sheets, *extra, template=None):
        p = self.sb.tmp / name
        fill(p, sheets)
        return self.sb.run("ingest.py", str(p), *extra)

    def table(self, n):
        return pd.read_csv(self.sb.data / f"{n}.csv", dtype=str, keep_default_na=False)


class IngestTests(IngestBase):
    def test_dry_run_writes_nothing(self):
        r = self.ingest("a.xlsx", {"NEW_PUBLICATIONS": [self.pub], "INGESTION": [self.row(value=50)]})
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertEqual(len(self.table("observations")), 0)

    def test_apply_then_reimport_is_rejected(self):
        sheets = {"NEW_PUBLICATIONS": [self.pub],
                  "INGESTION": [self.row(value=50, uncertainty_type="1_sigma", uncertainty_minus=2, uncertainty_plus=3, sample_count=12,
                                         measurement_start="2019", measurement_end="2019-08", spatial_resolution="30 m",
                                         sensitivity_threshold=0.1, sensitivity_threshold_unit="m", quality_flag="high", quality_reason="two methods agree",
                                         dataset_id="Supp. S1", source_row_id="7")]}
        r = self.ingest("a.xlsx", sheets, "--apply", "--submitter", "Tester (0000)")
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        o = self.table("observations")
        self.assertEqual(len(o), 1)
        row = o.iloc[0]
        self.assertTrue(row.obs_id.startswith("O") and len(row.obs_id) == 13)
        self.assertEqual((row.unc_type, row.unc_minus, row.unc_plus, row.n, row.submitted_by), ("1sigma", "2", "3", "12", "Tester (0000)"))
        self.assertEqual((row.measurement_start, row.measurement_end, row.spatial_resolution, row.sensitivity_threshold_unit, row.quality),
                         ("2019", "2019-08", "30 m", "m", "high"))
        self.assertEqual((row.dataset_id, row.source_row_id, row.pub_ids), ("Supp. S1", "7", "testa_2030"))
        r2 = self.ingest("b.xlsx", sheets, "--apply", "--submitter", "Tester (0000)")
        self.assertEqual(r2.returncode, 1)
        self.assertEqual(len(self.table("observations")), 1)

    def test_errors_block_the_whole_import(self):
        r = self.ingest("a.xlsx", {"NEW_PUBLICATIONS": [self.pub], "INGESTION": [self.row(value=50), self.row(metric_id="nonexistent", value=1)]},
                        "--apply", "--submitter", "T")
        self.assertEqual(r.returncode, 1)
        self.assertEqual(len(self.table("observations")), 0)
        self.assertEqual(len(self.table("sources")), 23)

    def test_missing_value_is_never_zero(self):
        r = self.ingest("a.xlsx", {"NEW_PUBLICATIONS": [self.pub], "INGESTION": [self.row()]})
        self.assertEqual(r.returncode, 1)
        self.assertIn("no value", r.stdout)

    def test_range_qualifier_and_category_kept_as_published(self):
        sheets = {"NEW_PUBLICATIONS": [self.pub], "INGESTION": [
            self.row(value_range_low=40, value_range_high=60),
            self.row(statistic_type="maximum", value=14, qualifier="<"),
            self.row(metric_id="earthquake_mechanism", statistic_type="category", value="Strike-slip")]}
        r = self.ingest("a.xlsx", sheets, "--apply", "--submitter", "T")
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        o = self.table("observations")
        a = o[o.metric_id == "mapped_surface_rupture_length"].set_index("statistic")
        self.assertEqual((a.loc["individual", "value_num"], a.loc["individual", "value_lo"], a.loc["individual", "value_hi"]), ("", "40", "60"))
        self.assertEqual((a.loc["maximum", "value_num"], a.loc["maximum", "qualifier"]), ("14", "<"))
        self.assertEqual(o[o.statistic == "category"].value_text.iloc[0], "Strike-slip")

    def test_comma_decimal_is_accepted(self):
        r = self.ingest("a.xlsx", {"NEW_PUBLICATIONS": [self.pub], "INGESTION": [self.row(value="3,2")]}, "--apply", "--submitter", "T")
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertEqual(self.table("observations").value_num.iloc[0], "3.2")

    def test_uncertainty_requires_type_and_positive_offsets(self):
        r = self.ingest("a.xlsx", {"NEW_PUBLICATIONS": [self.pub], "INGESTION": [self.row(value=5, uncertainty_minus=1)]})
        self.assertEqual(r.returncode, 1)
        r = self.ingest("b.xlsx", {"NEW_PUBLICATIONS": [self.pub], "INGESTION": [self.row(value=5, uncertainty_type="1_sigma", uncertainty_minus=-1)]})
        self.assertEqual(r.returncode, 1)
        r = self.ingest("c.xlsx", {"NEW_PUBLICATIONS": [self.pub], "INGESTION": [self.row(value=5, uncertainty_type="qualitative")]})
        self.assertEqual(r.returncode, 1)
        r = self.ingest("zero.xlsx", {"NEW_PUBLICATIONS": [self.pub], "INGESTION": [self.row(value=5, uncertainty_type="1_sigma", uncertainty_minus=0)]})
        self.assertEqual(r.returncode, 1)

    def test_dropdown_free_values_warn_without_blocking(self):
        wb = openpyxl.load_workbook(TEMPLATE)
        dropdown = next(dv for dv in wb["INGESTION"].data_validations.dataValidation if dv.type == "list")
        self.assertTrue(dropdown.showErrorMessage)
        self.assertEqual(dropdown.errorStyle, "warning")
        self.assertIn("Keep the free value", dropdown.error)

    def test_digitization_must_be_declared_with_uncertainty(self):
        r = self.ingest("a.xlsx", {"NEW_PUBLICATIONS": [self.pub], "INGESTION": [self.row(value=5, value_origin="digitized")]})
        self.assertEqual(r.returncode, 1)
        ok = self.row(value=5, value_origin="digitized", digitization_method="WebPlotDigitizer, Fig. 3a", digitization_unc=0.5)
        r = self.ingest("b.xlsx", {"NEW_PUBLICATIONS": [self.pub], "INGESTION": [ok]})
        self.assertEqual(r.returncode, 0, r.stdout)

    def test_database_derived_origin_is_refused(self):
        r = self.ingest("a.xlsx", {"NEW_PUBLICATIONS": [self.pub], "INGESTION": [self.row(metric_id="ssd_fraction", value=30, value_origin="database_derived")]})
        self.assertEqual(r.returncode, 1)
        self.assertIn("not allowed", r.stdout)

    def test_location_and_date_rules(self):
        bad = [self.row(value=1, longitude=-117.6, latitude=35.7),                                   # coordinates without origin
               self.row(value=2, longitude=-117.6, latitude=35.7, coordinate_origin="converted"),    # converted without method
               self.row(value=3, measurement_start="2019-13-45x"),                                    # malformed date
               self.row(value=4, sensitivity_threshold=0.1)]                                         # threshold without unit
        r = self.ingest("a.xlsx", {"NEW_PUBLICATIONS": [self.pub], "INGESTION": bad})
        self.assertEqual(r.returncode, 1)
        for code in ("coordinate_origin", "coordinate_conversion_method", "measurement_start", "sensitivity_threshold_unit"):
            self.assertIn(code, r.stdout, code)

    def test_duplicate_publication_by_doi_is_rejected(self):
        so = self.table("sources")
        r = self.ingest("a.xlsx", {"NEW_PUBLICATIONS": [dict(self.pub, doi=so[so.doi != ""].doi.iloc[0])]})
        self.assertEqual(r.returncode, 1)
        self.assertIn("DOI already registered", r.stdout)

    def test_duplicate_event_under_other_name_is_rejected(self):
        e = dict(event_name="Ridgecrest 2", event_date="2019-07-06", mw=7.1, latitude=35.77, longitude=-117.6, mechanism="Strike-slip",
                 coordinate_source="USGS")
        r = self.ingest("a.xlsx", {"NEW_EVENTS": [e]})
        self.assertEqual(r.returncode, 1)
        self.assertIn("possible duplicate", r.stdout)

    def test_event_resolved_by_alias_and_dropdown_format(self):
        al = pd.read_csv(self.sb.data / "event_aliases.csv", dtype=str, keep_default_na=False).iloc[0]
        rows = [self.row(target_id=al.alias, value=1), self.row(target_id=f"Whatever name ({'2019-07-06'}) [{self.ev}]", value=2),
                self.row(source_id="10.1000/test.2030.1", value=3)]
        r = self.ingest("a.xlsx", {"NEW_PUBLICATIONS": [self.pub], "INGESTION": rows}, "--apply", "--submitter", "T")
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertEqual(set(self.table("observations").event_id), {al.event_id, self.ev})

    def test_unknown_or_misspelled_column_is_reported(self):
        p = self.sb.tmp / "a.xlsx"
        fill(p, {"NEW_PUBLICATIONS": [self.pub], "INGESTION": [self.row(value=1)]})
        wb = openpyxl.load_workbook(p)
        wb["INGESTION"]["AY2"] = "valeu"
        wb.save(p)
        r = self.sb.run("ingest.py", str(p))
        self.assertEqual(r.returncode, 1)
        self.assertIn("unknown column 'valeu'", r.stdout)

    def test_excel_date_cell_in_new_event(self):
        e = dict(event_name="Newtown", event_date=dt.datetime(2031, 3, 4), mw=6.5, latitude=10, longitude=20, mechanism="Normal", coordinate_source="paper")
        r = self.ingest("a.xlsx", {"NEW_EVENTS": [e]}, "--apply", "--submitter", "T")
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertIn("2031-03-04", (self.sb.data / "events.csv").read_text())

    def test_excel_date_cells_in_measurement_dates(self):
        row = self.row(value=3.2, measurement_start=dt.date(2020, 1, 2), measurement_end=dt.datetime(2020, 2, 3))
        result = self.ingest("measurement_dates.xlsx", {"NEW_PUBLICATIONS": [self.pub], "INGESTION": [row]},
                             "--apply", "--submitter", "Tester")
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        observations = self.table("observations")
        self.assertEqual((observations.measurement_start.iloc[0], observations.measurement_end.iloc[0]),
                         ("2020-01-02", "2020-02-03"))

    def test_source_id_resolves_doi_pub_id_and_dropdown_label(self):
        publication = dict(citation_short="Testa et al. (2030)", citation_full=self.pub["citation_full"],
                           doi="10.1000/test.2030.1")
        tokens = ["https://doi.org/10.1000/test.2030.1", "testa_2030", "Testa et al. (2030) [testa_2030]"]
        for token in tokens:
            with self.subTest(source_id=token):
                result = self.ingest("source_id.xlsx", {"NEW_PUBLICATIONS": [publication],
                                        "INGESTION": [self.row(source_id=token, value=3.2)]})
                self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    def test_blank_trailing_and_hidden_rows_are_handled(self):
        path = self.sb.tmp / "hidden_rows.xlsx"
        fill(path, {"NEW_PUBLICATIONS": [self.pub], "INGESTION": [self.row(value=3.2), self.row(value=4.2)]})
        wb = openpyxl.load_workbook(path)
        ws = wb["INGESTION"]
        ws.row_dimensions[4].hidden = True
        ws.auto_filter.ref = f"A2:{ws.cell(2, ws.max_column).coordinate[0]}4"
        ws.append([None] * ws.max_column)
        ws.append([None] * ws.max_column)
        wb.save(path)
        result = self.sb.run("ingest.py", str(path))
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertIn("observations to add: 2", result.stdout)

    def test_500_row_series_imports_as_one_group(self):
        rows = [self.row(metric_id="surface_fault_slip", value=1 + i / 1000, component="strike-slip",
                         observation_group_id="large_series", source_location="Supp. Table S1", source_row_id=i + 1)
                for i in range(500)]
        result = self.ingest("large_series.xlsx", {"NEW_PUBLICATIONS": [self.pub], "INGESTION": rows},
                             "--apply", "--submitter", "Tester")
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertEqual(len(self.table("observations")), 0)
        arrays, values = self.table("arrays"), self.table("array_values")
        self.assertEqual((len(arrays), arrays.n_values.iloc[0], len(values)), (1, "500", 500))

    def test_contributor_branch_rows_merge_with_stable_distinct_hash_ids(self):
        contributors = [Sandbox("live"), Sandbox("live")]
        merged = Sandbox("live")
        try:
            for sb, value in zip(contributors, (3.2, 4.2)):
                workbook = sb.tmp / "branch.xlsx"
                fill(workbook, {"INGESTION": [self.row(source_id="antoine_2021", value=value)]})
                result = sb.run("ingest.py", str(workbook), "--apply", "--submitter", "Contributor")
                self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            rows = [pd.read_csv(sb.data / "observations.csv", dtype=str, keep_default_na=False).iloc[-1]
                    for sb in contributors]
            self.assertNotEqual(rows[0].obs_id, rows[1].obs_id)
            combined = pd.concat([pd.read_csv(sb.data / "observations.csv", dtype=str, keep_default_na=False)
                                  for sb in contributors], ignore_index=True)
            combined.to_csv(merged.data / "observations.csv", index=False, lineterminator="\n")
            result = merged.run("validate.py", "--quiet")
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        finally:
            for sb in contributors + [merged]:
                sb.cleanup()

    def test_apply_rolls_back_if_validation_fails_after_merge(self):
        workbook = self.sb.tmp / "rollback.xlsx"
        fill(workbook, {"INGESTION": [self.row(source_id="antoine_2021", value=3.2)]})
        before = {name: (self.sb.data / f"{name}.csv").read_bytes() for name in schema.FILES}
        redirected = {name: self.sb.data / f"{name}.csv" for name in schema.FILES}
        with patch.object(schema, "FILES", redirected), patch.object(schema, "REPORTS", self.sb.reports), \
             patch.object(sys, "argv", ["ingest.py", str(workbook), "--apply", "--submitter", "Tester"]), \
             patch.object(ingest.subprocess, "run", return_value=Mock(returncode=1)):
            with self.assertRaisesRegex(SystemExit, "import rolled back"):
                ingest.main()
        after = {name: (self.sb.data / f"{name}.csv").read_bytes() for name in schema.FILES}
        self.assertEqual(after, before)


class SeriesTests(IngestBase):
    def series(self, n=4, **kw):
        return [self.row(metric_id="surface_fault_slip", statistic_type="individual", value=v, component="strike-slip", observation_group_id="s1",
                         source_location="Supp. S1", source_spatial_data=5 * i, source_spatial_unit="km", **kw)
                for i, v in enumerate([0.8, 1.4, 2.1, 1.7][:n])]

    def test_group_of_individual_rows_becomes_one_array_with_computed_stats(self):
        r = self.ingest("a.xlsx", {"NEW_PUBLICATIONS": [self.pub], "INGESTION": self.series()}, "--apply", "--submitter", "T")
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertEqual(len(self.table("observations")), 0)
        a, v = self.table("arrays"), self.table("array_values")
        self.assertEqual((len(a), a.n_values.iloc[0], a.group_id.iloc[0], a.derive_stats.iloc[0]), (1, "4", "s1", "true"))
        self.assertEqual(list(v.value), ["0.8", "1.4", "2.1", "1.7"])
        self.assertEqual(list(v.source_position), ["0", "5", "10", "15"])
        b = self.sb.run("build.py")
        self.assertEqual(b.returncode, 0, b.stdout + b.stderr)
        self.assertIn('"computed_from_arrays": 4', b.stdout)

    def test_series_rows_must_agree(self):
        rows = self.series()
        rows[2]["component"] = "net slip"
        r = self.ingest("a.xlsx", {"NEW_PUBLICATIONS": [self.pub], "INGESTION": rows})
        self.assertEqual(r.returncode, 1)
        self.assertIn("differ in component", r.stdout)

    def test_series_rows_need_plain_numbers(self):
        rows = self.series()
        rows[1]["qualifier"] = "<"
        self.assertEqual(self.ingest("a.xlsx", {"NEW_PUBLICATIONS": [self.pub], "INGESTION": rows}).returncode, 1)

    def test_series_with_positions_needs_coordinate_origin_and_stores_them(self):
        rows = self.series()
        for i, r in enumerate(rows):
            r.update(longitude=-117.6 + 0.01 * i, latitude=35.7)
        self.assertEqual(self.ingest("a.xlsx", {"NEW_PUBLICATIONS": [self.pub], "INGESTION": rows}).returncode, 1)
        rows[0].update(coordinate_origin="source")
        r = self.ingest("b.xlsx", {"NEW_PUBLICATIONS": [self.pub], "INGESTION": rows}, "--apply", "--submitter", "T")
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertEqual(self.table("array_values").lat.iloc[0], "35.7")

    def test_group_with_summary_statistics_keeps_them_as_observations(self):
        rows = self.series() + [self.row(metric_id="surface_fault_slip", statistic_type="maximum", value=2.5, component="strike-slip",
                                         observation_group_id="s1", source_location="Supp. S1")]
        r = self.ingest("a.xlsx", {"NEW_PUBLICATIONS": [self.pub], "INGESTION": rows}, "--apply", "--submitter", "T")
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertEqual((len(self.table("observations")), len(self.table("arrays"))), (1, 1))
        self.assertEqual(self.table("observations").group_id.iloc[0], "s1")

    def test_single_row_with_group_id_is_a_plain_observation(self):
        r = self.ingest("a.xlsx", {"NEW_PUBLICATIONS": [self.pub], "INGESTION": self.series(1)}, "--apply", "--submitter", "T")
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertEqual((len(self.table("observations")), len(self.table("arrays"))), (1, 0))


class ExampleFileTests(IngestBase):
    def test_template_has_one_ingestion_sheet_with_spec_fields(self):
        wb = openpyxl.load_workbook(TEMPLATE)
        self.assertEqual(wb.sheetnames, ["INGESTION", "NEW_EVENTS", "NEW_PUBLICATIONS", "EVENT_ALIASES", "FIELDS", "HOW TO", "LISTS"])
        self.assertEqual([cell.value for cell in wb["INGESTION"][2]], spec.INGESTION_COLUMNS)
        self.assertFalse({"OBSERVATIONS", "ARRAYS", "ARRAY_VALUES"} & set(wb.sheetnames))

    def test_field_guide_covers_all_ingestion_columns(self):
        guide = (ROOT / "docs" / "INGESTION_FIELDS.md").read_text(encoding="utf8")
        for column in spec.INGESTION_COLUMNS:
            with self.subTest(column=column):
                self.assertIn(f"`{column}`", guide)

    def test_filled_example_cannot_be_imported(self):
        r = self.sb.run("ingest.py", str(EXAMPLE), "--apply", "--submitter", "T")
        self.assertEqual(r.returncode, 1)
        self.assertIn("EXAMPLE", r.stdout)
        self.assertEqual(len(self.table("observations")), 0)

    def test_blank_template_imports_nothing_and_is_clean(self):
        r = self.sb.run("ingest.py", str(TEMPLATE))
        self.assertEqual(r.returncode, 0, r.stdout)

    def test_example_has_every_pattern_the_protocol_describes(self):
        ws = openpyxl.load_workbook(EXAMPLE)["INGESTION"]
        head = [c.value for c in ws[2]]
        rows = [dict(zip(head, [c.value for c in r])) for r in ws.iter_rows(min_row=3)]
        used = {k for r in rows for k, v in r.items() if v not in (None, "")}
        for needed in ("value_range_low", "qualifier", "value_origin", "uncertainty_minus", "observation_group_id", "coordinate_conversion_method"):
            self.assertIn(needed, used)


if __name__ == "__main__":
    unittest.main()
