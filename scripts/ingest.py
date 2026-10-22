#!/usr/bin/env python3
"""Validate a contributor workbook and (optionally) import it into data/.

  python scripts/ingest.py my_file.xlsx                 dry run: report only, nothing is written
  python scripts/ingest.py my_file.xlsx --apply --submitter "Name (ORCID)"

Every problem is reported with the sheet and row number.  --apply is all-or-nothing: if there is
a single error nothing is written, and if the merged tables fail validation the previous files are
restored.  The pull request is the review step; reviewers can set review_status/reviewed_by with
scripts/review.py.
"""
import argparse
import datetime as dt
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
import schema                                                     # noqa: E402
import workbook_spec as ws_spec                                   # noqa: E402
import validate                                                   # noqa: E402
from citations import parse_label                                 # noqa: E402
from common import fmt_num, make_array_id, make_obs_id, read_table, to_float, write_table  # noqa: E402

SHEETS = {"INGESTION": ws_spec.INGESTION_COLUMNS, **ws_spec.SIDE_COLUMNS}


class Report:
    def __init__(self):
        self.errors, self.warnings = [], []

    def err(self, sheet, row, msg):
        self.errors.append((sheet, row, msg))

    def warn(self, sheet, row, msg):
        self.warnings.append((sheet, row, msg))


def cell(v):
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, str):
        return v.strip()
    return v


def read_sheet(wb, name, rep):
    if name not in wb.sheetnames:
        if name == "INGESTION":
            rep.err(name, "", "sheet INGESTION not found. Use templates/ingestion_template.xlsx (regenerate: python scripts/make_template.py)")
        return []
    ws = wb[name]
    allowed = set(SHEETS[name])
    header, hrow, rows = None, None, []
    for i, raw in enumerate(ws.iter_rows(values_only=True), start=1):
        vals = [cell(v) for v in raw]
        if header is None:
            names = [str(v).strip().lower() for v in vals]
            if sum(n in allowed for n in names) >= max(2, len(allowed) // 3):
                header, hrow = names, i
                for n in names:
                    if n and n not in allowed:
                        rep.err(name, i, f"unknown column '{n}'. Regenerate the template: python scripts/make_template.py")
                for n in allowed - set(names):
                    if n in ws_spec_required(name):
                        rep.err(name, i, f"required column '{n}' is missing")
            continue
        if all(v == "" for v in vals):
            continue
        if next((str(v) for v in vals if v != ""), "").startswith("#"):
            continue
        rows.append((i, {h: v for h, v in zip(header, vals) if h}))
    if header is None:
        rep.err(name, "", f"header row not found in sheet {name}: use the template")
    return rows


def ws_spec_required(name):
    fs = [f for _, g in ws_spec.FIELDS for f in g] if name == "INGESTION" else ws_spec.SIDE_SHEETS[name]
    return {f[0] for f in fs if f[1] == "yes"}


def num(v, rep, sheet, row, field, minimum=None, maximum=None, positive=False):
    if v == "":
        return None
    try:
        f = float(str(v).replace(",", ".")) if isinstance(v, str) else float(v)
    except ValueError:
        rep.err(sheet, row, f"{field}: '{v}' is not a number (no units, no ranges, no '±' - use the dedicated columns)")
        return None
    if minimum is not None and f < minimum:
        rep.err(sheet, row, f"{field}: {f:g} is below {minimum:g}")
    if maximum is not None and f > maximum:
        rep.err(sheet, row, f"{field}: {f:g} is above {maximum:g}")
    if positive and f <= 0:
        rep.err(sheet, row, f"{field}: {f:g} must be positive")
    return f


def norm_doi(d):
    d = str(d).strip()
    d = re.sub(r"^(https?://)?(dx\.)?doi\.org/", "", d, flags=re.I)
    return d


def slug(name):
    return re.sub(r"[^a-z0-9]", "", name.lower())


class Registry:
    """Lookup of existing + newly submitted events and publications."""

    def __init__(self, tabs):
        self.ev = tabs["events"].set_index("event_id").to_dict("index")
        self.alias = dict(zip(tabs["event_aliases"].alias, tabs["event_aliases"].event_id))
        self.src = tabs["sources"].set_index("pub_id").to_dict("index")
        self.doi = {r["doi"].lower(): p for p, r in self.src.items() if r["doi"]}
        self.metrics = tabs["metrics"].set_index("metric_id").to_dict("index")

    def event(self, token):
        t = str(token).strip()
        m = re.search(r"\[([^\[\]]+)\]\s*$", t)
        if m and m.group(1) in self.ev:
            return m.group(1)
        if t in self.ev:
            return t
        if t in self.alias:
            return self.alias[t]
        names = [e for e, r in self.ev.items() if r["event_name"].lower() == t.lower()]
        if len(names) == 1:
            return names[0]
        for a, e in self.alias.items():
            if a.lower() == t.lower():
                return e
        return None

    def pub(self, token):
        t = str(token).strip()
        m = re.search(r"\[([^\[\]]+)\]\s*$", t)
        if m and m.group(1) in self.src:
            return m.group(1)
        if t in self.src:
            return t
        d = norm_doi(t).lower()
        if d in self.doi:
            return self.doi[d]
        pubs, ann = parse_label(t)
        if len(pubs) == 1 and pubs[0]["pub_id"] in self.src and not ann:
            return pubs[0]["pub_id"]
        return None


def process_events(rows, reg, rep, new):
    for i, r in rows:
        S = "NEW_EVENTS"
        for f in ("event_name", "event_date", "mw", "latitude", "longitude", "mechanism", "coordinate_source"):
            if r.get(f, "") == "":
                rep.err(S, i, f"{f} is required")
        name, date = str(r.get("event_name", "")), str(r.get("event_date", ""))
        if date and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", date):
            rep.err(S, i, f"event_date '{date}' must be YYYY-MM-DD")
            continue
        mw = num(r.get("mw", ""), rep, S, i, "mw", 3, 10)
        lat = num(r.get("latitude", ""), rep, S, i, "latitude", -90, 90)
        lon = num(r.get("longitude", ""), rep, S, i, "longitude", -180, 180)
        hd = num(r.get("hypocenter_depth_km", ""), rep, S, i, "hypocenter_depth_km", 0, 100)
        cd = num(r.get("centroid_depth_km", ""), rep, S, i, "centroid_depth_km", 0, 100)
        if r.get("mechanism", "") and r["mechanism"] not in schema.MECHANISMS:
            rep.err(S, i, f"mechanism must be one of {schema.MECHANISMS}")
        if not name or not date:
            continue
        eid = f"event:{slug(name)}{date[:4]}"
        if eid in reg.ev or eid in new["events"]:
            rep.err(S, i, f"event id {eid} already exists - use the existing event (add an alias if the name differs)")
            continue
        if r.get("confirm_distinct", "").lower() != "yes":
            d0 = dt.date.fromisoformat(date)
            for oid, o in list(reg.ev.items()) + list(new["events"].items()):
                if not o["event_date"]:
                    continue
                try:
                    d1 = dt.date.fromisoformat(o["event_date"])
                except ValueError:
                    continue
                if abs((d0 - d1).days) <= 2 and slug(o["event_name"])[:6] == slug(name)[:6]:
                    rep.err(S, i, f"possible duplicate of {oid} ({o['event_name']}, {o['event_date']}). Add an alias instead, "
                                  "or set confirm_distinct=yes if this really is a different earthquake (foreshock/aftershock/other).")
        row = dict(event_id=eid, event_name=name, event_date=date, mw=fmt_num(mw), latitude=fmt_num(lat), longitude=fmt_num(lon),
                   hypocenter_depth_km=fmt_num(hd), centroid_depth_km=fmt_num(cd), mechanism=r.get("mechanism", ""),
                   region=r.get("region", ""), coordinate_source=r.get("coordinate_source", ""),
                   sequence_id=r.get("sequence_id", ""), notes=r.get("notes", ""))
        new["events"][eid] = row
        reg.ev[eid] = row
        new["row_of"][eid] = (S, i)
        for a in filter(None, (x.strip() for x in str(r.get("aliases", "")).split("|"))):
            if a in reg.alias or a in reg.ev:
                rep.err(S, i, f"alias '{a}' already exists")
            else:
                reg.alias[a] = eid
                new["aliases"].append(dict(alias=a, alias_type="name_variant", event_id=eid, reason="added with the event"))


def process_pubs(rows, reg, rep, new):
    S = "NEW_PUBLICATIONS"
    for i, r in rows:
        cs = str(r.get("citation_short", ""))
        if ws_spec.EXAMPLE_MARK in str(r.get("note", "")):
            rep.err(S, i, "this is an EXAMPLE publication; the example file cannot be imported")
        pubs, ann = parse_label(cs)
        if len(pubs) != 1 or ann:
            rep.err(S, i, f"citation_short '{cs}' must be exactly one citation like 'Surname et al. (2020)' (found {len(pubs)}{', stray text: ' + ann if ann else ''})")
            continue
        p = pubs[0]
        kind = r.get("kind", "") or ("agency" if p["kind"] == "agency" else "publication")
        if kind not in schema.KINDS:
            rep.err(S, i, f"kind must be one of {schema.KINDS}")
            continue
        full = str(r.get("citation_full", ""))
        if len(full) < 20:
            rep.err(S, i, "citation_full is required (full reference)")
        doi = norm_doi(r.get("doi", ""))
        if doi and not re.fullmatch(r"10\.\d{4,9}/\S+", doi):
            rep.err(S, i, f"doi '{doi}' is malformed (expected 10.xxxx/...)")
        if kind == "publication" and not doi and not r.get("no_doi_reason", ""):
            rep.err(S, i, "a publication needs a doi, or no_doi_reason")
        pid = p["pub_id"]
        if doi and doi.lower() in reg.doi:
            rep.err(S, i, f"DOI already registered as {reg.doi[doi.lower()]} - use it in 'publication' (spelling variants are not new publications)")
            continue
        if pid in reg.src:
            rep.err(S, i, f"pub_id {pid} already exists ({reg.src[pid]['citation_short']}). Same author and year but another paper? "
                          "add a letter: 'Surname et al. (2020b)'. Same paper? use the existing one.")
            continue
        row = dict(pub_id=pid, citation_short=p["short"], citation_full=full, doi=doi, kind=kind,
                   needs_full_reference="false", legacy_source_ids="",
                   note=r.get("note", "") or (f"no DOI: {r['no_doi_reason']}" if r.get("no_doi_reason") else ""))
        new["sources"][pid] = row
        reg.src[pid] = row
        if doi:
            reg.doi[doi.lower()] = pid


def process_aliases(rows, reg, rep, new):
    S = "EVENT_ALIASES"
    for i, r in rows:
        eid = reg.event(r.get("event", ""))
        a = str(r.get("alias", ""))
        if not eid:
            rep.err(S, i, f"event '{r.get('event', '')}' not found")
        elif not a:
            rep.err(S, i, "alias is empty")
        elif a in reg.alias or a in reg.ev:
            rep.err(S, i, f"alias '{a}' already exists")
        else:
            reg.alias[a] = eid
            new["aliases"].append(dict(alias=a, alias_type="name_variant", event_id=eid, reason=r.get("reason", "") or "added by contributor"))


UNC_ALIASES = {"1_sigma": "1sigma", "1sigma": "1sigma", "2_sigma": "2sigma", "2sigma": "2sigma", "95_ci": "ci95", "ci95": "ci95",
               "credible_interval": "credible_interval", "range": "minmax_range", "minmax_range": "minmax_range",
               "qualitative": "qualitative", "unspecified": "unspecified", "not_reported": "not_reported"}
ORIGIN_ALIASES = {"": "reported", "reported": "reported", "digitized": "figure_digitized", "figure_digitized": "figure_digitized"}
SERIES_SAME = ("event_id", "metric_id", "pub_ids", "compiled_in", "component", "segment", "model_id", "unit", "method",
               "source_location", "unc_type", "origin")
LOC_COLS = ["spatial_representation", "longitude", "latitude", "depth_m", "horizontal_location_uncertainty_m",
            "depth_location_uncertainty_m", "coordinate_origin", "coordinate_conversion_method", "source_spatial_data",
            "source_spatial_unit", "source_reference_description"]
CTX_COLS = ["spatial_resolution", "sensitivity_threshold", "sensitivity_threshold_unit", "measurement_start", "measurement_end",
            "quality_reason", "confidence_level"]


def unit_for(S, i, r, reg, rep, mid):
    m = reg.metrics.get(mid)
    if not m:
        return str(r.get("unit", ""))
    cu, u = m["canonical_unit"], str(r.get("unit", ""))
    if cu:
        if u and u != cu:
            rep.err(S, i, f"unit '{u}' differs from the canonical unit '{cu}': convert, and keep the printed form in reported_text")
        return cu
    if m["value_type"] == "numeric" and not u:
        rep.err(S, i, f"metric {mid} has no canonical unit: give the unit explicitly")
    return u


def parse_row(i, r, reg, rep):
    """One INGESTION row -> dict of canonical fields (errors go to the report)."""
    S = "INGESTION"
    p = {"_row": i}
    if schema_example(r):
        rep.err(S, i, "this is an EXAMPLE row (comments start with [EXAMPLE]); the example file cannot be imported")
    eid = reg.event(r.get("target_id", ""))
    if not eid:
        rep.err(S, i, f"target_id '{r.get('target_id', '')}' not found. Pick it from the list, or add it in NEW_EVENTS (other names: EVENT_ALIASES).")
    mid = str(r.get("metric_id", "")).strip()
    mm = re.search(r"\[([^\[\]]+)\]\s*$", mid)
    mid = mm.group(1) if mm else mid
    m = reg.metrics.get(mid)
    if not m:
        rep.err(S, i, f"metric_id '{mid}' is not in the metric catalogue (data/metrics.csv)")
    pubs = []
    for tok in re.split(r"[|;]", str(r.get("source_id", ""))):
        if tok.strip():
            q = reg.pub(tok)
            if q is None:
                rep.err(S, i, f"source_id '{tok.strip()}' not found. Add it in NEW_PUBLICATIONS (full reference + DOI).")
            else:
                pubs.append(q)
    if not pubs:
        rep.err(S, i, "source_id is required")
    comp = ""
    if r.get("compiled_in", ""):
        comp = reg.pub(r["compiled_in"]) or ""
        if not comp:
            rep.err(S, i, f"compiled_in '{r['compiled_in']}' not found in publications")
        elif comp in pubs:
            rep.err(S, i, "compiled_in is the same as source_id: leave compiled_in empty unless you read the value in a compilation")
    component = str(r.get("component", ""))
    if component not in schema.COMPONENTS:
        rep.err(S, i, f"component '{component}' not in the controlled list {schema.COMPONENTS[1:]}")
    if len(str(r.get("source_location", ""))) < 3:
        rep.err(S, i, "source_location is required (page / table / figure caption / supplement file)")
    if not r.get("method", ""):
        rep.err(S, i, "method is required (field, OIC, InSAR, GNSS, seismic, geologic_mapping, laboratory, joint_inversion, ...)")
    va = str(r.get("verified_against", "")).lower()
    if va not in schema.VERIFIED_AGAINST:
        rep.err(S, i, "verified_against must be empty, 'text' or 'table'")
    origin_in = str(r.get("value_origin", "")).lower()
    if origin_in == "database_derived":
        rep.err(S, i, "value_origin database_derived is not allowed: the database computes min/mean/median/max from series itself, and never derives OFD/SSD")
    origin = ORIGIN_ALIASES.get(origin_in)
    if origin is None:
        rep.err(S, i, "value_origin must be 'reported' or 'digitized'")
        origin = "reported"
    dm, du = str(r.get("digitization_method", "")), num(r.get("digitization_unc", ""), rep, S, i, "digitization_unc", positive=True)
    if origin == "figure_digitized":
        if not dm or not du:
            rep.err(S, i, "digitized needs digitization_method (software, axis calibration, panel) and a positive digitization_unc in data units")
    elif dm or du:
        rep.err(S, i, "digitization fields are only for value_origin=digitized")
    # ---- value ----
    cat = bool(m) and m["value_type"] == "categorical"
    stat = str(r.get("statistic_type", "")).lower()
    if stat == "percentile":
        rep.err(S, i, "statistic_type 'percentile' is not supported yet (no column for the percentile level)")
    if cat and stat == "":
        stat = "category"
    allowed = schema.CATEGORICAL_STATS if cat else schema.NUMERIC_STATS
    if m and stat not in allowed:
        rep.err(S, i, f"statistic_type '{stat}' not allowed for {mid} (allowed: {allowed})")
    raw = r.get("value", "")
    v = text = None
    if cat:
        text = str(raw)
        if not text:
            rep.err(S, i, "value (the category) is required for a categorical metric")
        if r.get("value_range_low", "") != "" or r.get("value_range_high", "") != "":
            rep.err(S, i, "categorical metric: no value range")
    else:
        v = num(raw, rep, S, i, "value")
    lo, hi = num(r.get("value_range_low", ""), rep, S, i, "value_range_low"), num(r.get("value_range_high", ""), rep, S, i, "value_range_high")
    q = str(r.get("qualifier", ""))
    if q not in schema.QUALIFIERS:
        rep.err(S, i, f"qualifier must be one of {schema.QUALIFIERS[1:]}")
    rt = ""
    if not cat:
        if v is None and lo is None and hi is None and raw == "":
            rep.err(S, i, "no value: give 'value', or value_range_low + value_range_high. Never write 0 for 'missing'.")
        if v is not None and (lo is not None or hi is not None):
            rep.err(S, i, "give either value or a range, not both (a +/- belongs in the uncertainty columns)")
        if (lo is None) != (hi is None):
            rep.err(S, i, "a range needs both value_range_low and value_range_high (use qualifier < or > for a one-sided bound)")
        if lo is not None and hi is not None:
            rt = "published_range"
            if lo > hi:
                rep.err(S, i, "value_range_low is larger than value_range_high")
        if q and v is None:
            rep.err(S, i, "qualifier needs a value")
        if m:
            pmin, pmax = to_float(m["plausible_min"]), to_float(m["plausible_max"])
            for x in (v, lo, hi):
                if x is not None and ((pmin is not None and x < pmin) or (pmax is not None and x > pmax)):
                    rep.warn(S, i, f"{mid} value {x:g} is outside the plausible range [{pmin}, {pmax}] - double-check the unit")
                    break
    unit = unit_for(S, i, r, reg, rep, mid) if (m and not cat) else ""
    # ---- uncertainty ----
    ut_in = str(r.get("uncertainty_type", "")).lower()
    ut = UNC_ALIASES.get(ut_in, None) if ut_in else ""
    if ut_in and ut is None:
        rep.err(S, i, f"uncertainty_type '{ut_in}' must be one of {sorted(UNC_ALIASES)}")
        ut = ""
    um = num(r.get("uncertainty_minus", ""), rep, S, i, "uncertainty_minus", positive=True)
    up = num(r.get("uncertainty_plus", ""), rep, S, i, "uncertainty_plus", positive=True)
    udesc = str(r.get("uncertainty_description", ""))
    if ut == "not_reported" and (um is not None or up is not None):
        rep.err(S, i, "uncertainty_type not_reported but offsets given")
    if (um is not None or up is not None) and not ut:
        rep.err(S, i, "uncertainty offsets given without uncertainty_type: say how the paper defines it ('unspecified' if it does not)")
    if ut in ("1sigma", "2sigma", "ci95", "credible_interval", "minmax_range", "unspecified") and um is None and up is None:
        rep.err(S, i, f"uncertainty_type {ut_in} needs uncertainty_minus and/or uncertainty_plus (positive offsets from the value)")
    if ut == "qualitative" and not udesc:
        rep.err(S, i, "qualitative uncertainty needs uncertainty_description")
    if (um is not None or up is not None) and v is None:
        rep.err(S, i, "uncertainty offsets need a central value (not a range-only value)")
    n = num(r.get("sample_count", ""), rep, S, i, "sample_count", 1)
    qual = str(r.get("quality_flag", "")) or "not_assessed"
    if qual not in schema.QUALITY:
        rep.err(S, i, f"quality_flag must be one of {schema.QUALITY}")
    sr = str(r.get("spatial_representation", ""))
    if sr and sr not in schema.SPATIAL_REPRESENTATIONS:
        rep.err(S, i, f"spatial_representation must be one of {schema.SPATIAL_REPRESENTATIONS}")
    dstat = str(r.get("derive_stats", "")).lower() or "yes"
    if dstat not in ("yes", "no"):
        rep.err(S, i, "derive_stats must be yes or no")
    p.update(
        event_id=eid or "", metric_id=mid, pub_ids="|".join(sorted(set(pubs))), compiled_in=comp, component=component,
        segment=str(r.get("segment", "")), model_id=str(r.get("model_id", "")), source_location=str(r.get("source_location", "")),
        method=str(r.get("method", "")), method_details=str(r.get("method_details", "")), origin=origin, digitization_method=dm,
        digitization_unc=fmt_num(du), review_status=schema.VERIFIED_AGAINST.get(va, "draft"), statistic=stat, value=v, lo=lo, hi=hi,
        qualifier=q, range_type=rt, text=text, unit=unit, n=fmt_num(n), unc_type=ut or "not_reported", um=fmt_num(um), up=fmt_num(up),
        unc_note=udesc, reported_text=str(r.get("reported_text", "")), quality=qual, comment=str(r.get("comments", "")),
        group_id=str(r.get("observation_group_id", "")), dataset_id=str(r.get("dataset_id", "")), source_row_id=str(r.get("source_row_id", "")),
        derive="true" if dstat == "yes" else "false", cat=cat,
        loc={k: str(r.get(k, "")) for k in LOC_COLS}, ctx={k: str(r.get(k, "")) for k in CTX_COLS},
        confidence_level=str(r.get("confidence_level", "")))
    p["ctx"]["quality_reason"] = str(r.get("quality_reason", ""))
    p["ctx"]["confidence_level"] = str(r.get("confidence_level", ""))
    if p["loc"]["longitude"] or p["loc"]["latitude"]:
        for k in ("longitude", "latitude"):
            p["loc"][k] = fmt_num(num(r.get(k, ""), rep, S, i, k, -180 if k == "longitude" else -90, 180 if k == "longitude" else 90))
    return p


def schema_example(r):
    return ws_spec.EXAMPLE_MARK in str(r.get("comments", ""))


def make_observation(p):
    row = dict(
        obs_id="", event_id=p["event_id"], metric_id=p["metric_id"], pub_ids=p["pub_ids"], compiled_in=p["compiled_in"],
        statistic=p["statistic"], value_num=fmt_num(p["value"]), qualifier=p["qualifier"], value_lo=fmt_num(p["lo"]),
        value_hi=fmt_num(p["hi"]), range_type=p["range_type"], value_text=p["text"] or "", unit=p["unit"], component=p["component"],
        segment=p["segment"], model_id=p["model_id"], n=p["n"], unc_type=p["unc_type"], unc_minus=p["um"], unc_plus=p["up"],
        unc_note=p["unc_note"], reported_text=p["reported_text"], source_location=p["source_location"], method=p["method"],
        method_details=p["method_details"], origin=p["origin"], digitization_method=p["digitization_method"],
        digitization_unc=p["digitization_unc"], dataset_id=p["dataset_id"], group_id=p["group_id"], source_row_id=p["source_row_id"],
        review_status=p["review_status"], quality=p["quality"], flags="value_from_range" if p["lo"] is not None else "",
        comment=p["comment"], legacy_label="", submitted_by="", submitted_on="", reviewed_by="", reviewed_on="")
    row.update(p["loc"])
    row.update(p["ctx"])
    row["obs_id"] = make_obs_id(row)
    return row


def make_series(group, ps, rep, new):
    """ps: parsed rows (>= 2) of one observation_group_id, all statistic 'individual'."""
    S, first = "INGESTION", ps[0]
    ok = True
    for k in SERIES_SAME:
        vals = {p[k] for p in ps}
        if len(vals) > 1:
            rep.err(S, first["_row"], f"series '{group}': rows {[p['_row'] for p in ps]} differ in {k} ({sorted(map(str, vals))}). "
                                      "Rows of one series must share it; use another observation_group_id for the others.")
            ok = False
    for p in ps:
        if p["value"] is None or p["lo"] is not None or p["qualifier"] or p["cat"]:
            rep.err(S, p["_row"], f"series '{group}': every row needs a plain numeric value (no range, qualifier or category)")
            ok = False
        if p["loc"]["longitude"] and not first["loc"]["coordinate_origin"]:
            rep.err(S, p["_row"], "series positions need coordinate_origin on the first row of the series")
            ok = False
    anyunc = any(p["um"] or p["up"] for p in ps)
    if anyunc and first["unc_type"] == "not_reported":
        rep.err(S, first["_row"], f"series '{group}': offsets given but no uncertainty_type")
        ok = False
    if not ok:
        return
    row = dict(
        array_id="", event_id=first["event_id"], metric_id=first["metric_id"], pub_ids=first["pub_ids"], compiled_in=first["compiled_in"],
        component=first["component"], segment=first["segment"], model_id=first["model_id"], unit=first["unit"],
        source_location=first["source_location"], method=first["method"], method_details=first["method_details"],
        derive_stats=first["derive"], unc_type=first["unc_type"], unc_note=first["unc_note"], n_values=str(len(ps)),
        origin=first["origin"], digitization_method=first["digitization_method"], digitization_unc=first["digitization_unc"],
        dataset_id=first["dataset_id"], group_id=group, review_status=first["review_status"], quality=first["quality"], flags="",
        comment=first["comment"], legacy_array_id="", submitted_by="", submitted_on="", reviewed_by="", reviewed_on="")
    for k in ("spatial_representation", "coordinate_origin", "coordinate_conversion_method", "source_spatial_unit", "source_reference_description"):
        row[k] = first["loc"][k]
    row.update(first["ctx"])
    row["array_id"] = make_array_id(row)
    if row["array_id"] in new["arrays"]:
        rep.err(S, first["_row"], f"series '{group}' has the same event/metric/study/component/segment/model/source_location as another series in this file")
        return
    new["arrays"][row["array_id"]] = row
    new["row_of"][row["array_id"]] = (S, first["_row"])
    new["array_values"][row["array_id"]] = [
        dict(array_id=row["array_id"], idx=str(k), value=fmt_num(p["value"]), unc_minus=p["um"], unc_plus=p["up"],
             lat=p["loc"]["latitude"], lon=p["loc"]["longitude"], source_position=p["loc"]["source_spatial_data"])
        for k, p in enumerate(ps)]


def process_ingestion(rows, reg, rep, new):
    parsed = [parse_row(i, r, reg, rep) for i, r in rows]
    groups = {}
    for p in parsed:
        if p["group_id"]:
            groups.setdefault(p["group_id"], []).append(p)
    series_rows = set()
    for g, ps in groups.items():
        ind = [p for p in ps if p["statistic"] == "individual"]
        if len(ind) >= 2:
            make_series(g, ind, rep, new)
            series_rows |= {p["_row"] for p in ind}
    for p in parsed:
        if p["_row"] in series_rows:
            continue
        if p["loc"]["longitude"] and not p["loc"]["coordinate_origin"]:
            rep.err("INGESTION", p["_row"], "longitude/latitude given without coordinate_origin ('source' or 'converted')")
        row = make_observation(p)
        if row["obs_id"] in new["obs"]:
            rep.err("INGESTION", p["_row"], f"duplicate of row {new['row_of'][row['obs_id']][1]} of this workbook (same study, event, metric, statistic, component, segment, model and value)")
            continue
        new["obs"][row["obs_id"]] = row
        new["row_of"][row["obs_id"]] = ("INGESTION", p["_row"])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--submitter", default="")
    args = ap.parse_args()
    path = Path(args.workbook)
    rep = Report()
    wb = openpyxl.load_workbook(path, data_only=True)
    tabs = {n: read_table(n) for n in schema.FILES}
    reg = Registry(tabs)
    new = dict(events={}, sources={}, aliases=[], obs={}, arrays={}, array_values={}, row_of={})
    sheets = {n: read_sheet(wb, n, rep) for n in SHEETS}
    process_events(sheets["NEW_EVENTS"], reg, rep, new)
    process_pubs(sheets["NEW_PUBLICATIONS"], reg, rep, new)
    process_aliases(sheets["EVENT_ALIASES"], reg, rep, new)
    process_ingestion(sheets["INGESTION"], reg, rep, new)

    # merge in memory and run the full validator; report only issues that this submission introduces
    before = validate.collect(tabs).df()
    before_ids = set(before.code + "|" + before.key)
    merged = {k: v.copy() for k, v in tabs.items()}
    stamp = dt.date.today().isoformat()
    submitter = args.submitter or "(dry run)"

    def add(name, rows):
        if rows:
            merged[name] = pd.concat([merged[name], pd.DataFrame(rows, columns=schema.COLUMNS[name])], ignore_index=True).fillna("")
    for d in new["obs"].values():
        d.update(submitted_by=submitter, submitted_on=stamp)
    for d in new["arrays"].values():
        d.update(submitted_by=submitter, submitted_on=stamp)
    add("events", list(new["events"].values()))
    add("sources", list(new["sources"].values()))
    add("event_aliases", new["aliases"])
    already = set(merged["observations"].obs_id) & set(new["obs"])
    for oid in already:
        sh, rw = new["row_of"][oid]
        rep.err(sh, rw, "identical row already exists in the database (same study, event, metric, statistic, component, segment, model, value)")
    add("observations", [d for o, d in new["obs"].items() if o not in already])
    for aid in set(merged["arrays"].array_id) & set(new["arrays"]):
        sh, rw = new["row_of"][aid]
        rep.err(sh, rw, "an array with the same event/metric/publication/component/segment/model/source_location already exists")
    add("arrays", [d for a, d in new["arrays"].items() if a not in set(merged["arrays"].array_id)])
    add("array_values", [r for a, rows in new["array_values"].items() for r in rows])
    after = validate.collect(merged).df()
    for _, r in after.iterrows():
        if f"{r.code}|{r.key}" in before_ids:
            continue
        where = ""
        for oid, (sh, rw) in new["row_of"].items():
            if oid in r.key or oid in r.detail:
                where = f"{sh} row {rw}"
                break
        (rep.errors if r.severity == "ERROR" else rep.warnings).append((where or "database", "", f"[{r.code}] {r.detail}"))

    # ---- report ----
    lines = [f"# Ingestion report: {path.name}", "",
             f"- new events: {len(new['events'])}   new publications: {len(new['sources'])}   new aliases: {len(new['aliases'])}",
             f"- observations to add: {len(new['obs'])}   arrays: {len(new['arrays'])} "
             f"({sum(len(v) for v in new['array_values'].values())} values)",
             f"- errors: {len(rep.errors)}   warnings: {len(rep.warnings)}", ""]
    if rep.errors:
        lines += ["## Errors (nothing is imported until these are fixed)", ""]
        lines += [f"- **{s}{' row ' + str(r) if r != '' else ''}**: {m}" for s, r, m in rep.errors]
    if rep.warnings:
        lines += ["", "## Warnings", ""] + [f"- {s}{' row ' + str(r) if r != '' else ''}: {m}" for s, r, m in rep.warnings]
    if not rep.errors:
        lines += ["", "No errors. " + ("Imported." if args.apply else "Re-run with --apply --submitter 'Name (ORCID)' to import.")]
    schema.REPORTS.mkdir(exist_ok=True)
    out = schema.REPORTS / f"ingest_{path.stem}.md"
    out.write_text("\n".join(lines) + "\n", encoding="utf8")
    print("\n".join(lines))
    print(f"\n(report: {out})")
    if rep.errors:
        sys.exit(1)
    if not args.apply:
        return
    if not args.submitter:
        sys.exit("--apply needs --submitter 'Name (ORCID)'")
    backup = Path(tempfile.mkdtemp(prefix="ingest_backup_"))
    for n, f in schema.FILES.items():
        shutil.copy(f, backup / f.name)
    try:
        for n, df in merged.items():
            write_table(n, df)
        res = subprocess.run([sys.executable, "-I", str(Path(__file__).parent / "validate.py"), "--quiet"])
        if res.returncode != 0:
            raise RuntimeError("validation of the merged tables failed")
    except Exception as e:                                   # restore everything
        for n, f in schema.FILES.items():
            shutil.copy(backup / f.name, f)
        shutil.rmtree(backup, ignore_errors=True)
        sys.exit(f"import rolled back: {e}")
    shutil.rmtree(backup, ignore_errors=True)
    print("imported. Next: python scripts/build.py && git add data && git commit && open a pull request")


if __name__ == "__main__":
    main()
